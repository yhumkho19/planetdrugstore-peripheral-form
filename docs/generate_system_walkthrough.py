from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = "event-scheduling-stock-ordering.pptx"
W, H = 13.333, 7.5
FONT = "Aptos"

INK = "202A35"
MUTED = "65717D"
PAPER = "F7F5F1"
WHITE = "FFFFFF"
MAROON = "76232E"
RED = "C94750"
TEAL = "248B7D"
MINT = "E4F2ED"
GOLD = "D79A2B"
AMBER = "FFF1D5"
LINE = "D9DDE0"
PALE_RED = "F9E9E8"
PALE_BLUE = "EAF0F4"
GREEN = "2F7A53"


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background() if not line else None
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except (IndexError, AttributeError):
            pass
    return shape


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04,
         font=FONT, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(2)
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.italic = italic
        paragraph.font.color.rgb = rgb(color)
    return box


def line(slide, x1, y1, x2, y2, color=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def pill(slide, label, x, y, w, fill, color=INK, size=10):
    rect(slide, x, y, w, 0.3, fill, radius=True)
    text(slide, label, x + 0.04, y + 0.01, w - 0.08, 0.26,
         size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE, margin=0.01)


def card(slide, x, y, w, h, title, body, accent=MAROON,
         fill=WHITE, title_size=16, body_size=12.2):
    rect(slide, x, y, w, h, fill, radius=True, line=LINE)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, title, x + 0.22, y + 0.18, w - 0.4, 0.42,
         size=title_size, color=INK, bold=True)
    text(slide, body, x + 0.22, y + 0.68, w - 0.42, h - 0.82,
         size=body_size, color=MUTED)


def footer(slide, number, label="PLANET DRUGSTORE  /  OPERATIONS WALKTHROUGH"):
    line(slide, 0.55, 7.08, 12.78, 7.08, LINE, 0.8)
    text(slide, label, 0.58, 7.14, 10.8, 0.2, size=8.5, color=MUTED, bold=True)
    text(slide, f"{number:02d}", 12.18, 7.12, 0.55, 0.22,
         size=9, color=MAROON, bold=True, align=PP_ALIGN.RIGHT)


def base_slide(prs, bg=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def header(slide, eyebrow, title, subtitle=""):
    text(slide, eyebrow.upper(), 0.62, 0.35, 11.8, 0.25,
         size=9.5, color=TEAL, bold=True)
    text(slide, title, 0.62, 0.7, 12.0, 0.55,
         size=27, color=INK, bold=True)
    if subtitle:
        text(slide, subtitle, 0.64, 1.35, 12.0, 0.46,
             size=12.3, color=MUTED)


def step(slide, number, title, body, x, y, w, color):
    rect(slide, x, y, w, 2.1, WHITE, radius=True, line=LINE)
    rect(slide, x + 0.2, y + 0.2, 0.42, 0.42, color, radius=True)
    text(slide, number, x + 0.2, y + 0.2, 0.42, 0.42,
         size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE, margin=0.01)
    text(slide, title, x + 0.2, y + 0.78, w - 0.4, 0.43,
         size=15, color=INK, bold=True)
    text(slide, body, x + 0.2, y + 1.28, w - 0.4, 0.65,
         size=11.2, color=MUTED)


def arrow(slide, x, y, w=0.34, color=GOLD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(0.28)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Event Scheduling & Stock Ordering"
    prs.core_properties.subject = "Operational walkthrough for Planet Drugstore"
    prs.core_properties.author = "Planet Drugstore"
    prs.core_properties.keywords = "event scheduling, deployment roster, stock request, SR#"

    # 1. Cover
    slide = base_slide(prs, INK)
    rect(slide, 0, 0, 0.18, H, RED)
    rect(slide, 8.8, 0.65, 3.9, 6.15, "263642", radius=True)
    text(slide, "OPERATIONS WALKTHROUGH", 0.82, 0.82, 6.6, 0.3,
         size=10, color="9AD3C5", bold=True)
    text(slide, "Event Scheduling\n& Stock Ordering", 0.8, 1.55, 7.6, 1.9,
         size=34, color=WHITE, bold=True)
    text(slide, "How teams publish event rosters, claim roles, submit SR#s,\nand manage cutoffs.",
         0.84, 3.85, 7.3, 0.82, size=16, color="D5DEE4")
    pill(slide, "EVENT ROSTER", 0.85, 5.35, 1.65, "D8EEE8", TEAL, 9)
    pill(slide, "STOCK REQUEST", 2.65, 5.35, 1.85, "F5E4E2", MAROON, 9)
    # Two connected workflow lanes
    text(slide, "EVENT SCHEDULING", 9.18, 1.05, 3.0, 0.25,
         size=10, color="9AD3C5", bold=True)
    for yy, label, color in [(1.55, "Prepare roster", TEAL),
                             (2.38, "Publish", RED),
                             (3.21, "Employees sign up", GOLD)]:
        rect(slide, 9.15, yy, 2.95, 0.53, "344652", radius=True)
        rect(slide, 9.34, yy + 0.15, 0.2, 0.2, color, radius=True)
        text(slide, label, 9.72, yy + 0.09, 2.15, 0.32,
             size=13, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    line(slide, 9.45, 2.08, 9.45, 2.31, "71838E", 1.2)
    line(slide, 9.45, 2.91, 9.45, 3.14, "71838E", 1.2)
    text(slide, "STOCK ORDERING", 9.18, 4.22, 3.0, 0.25,
         size=10, color="F0C779", bold=True)
    for yy, label, color in [(4.7, "Set days & cutoff", GOLD),
                             (5.53, "Submit SR#", TEAL),
                             (6.36, "Lock / reopen", RED)]:
        rect(slide, 9.15, yy, 2.95, 0.42, "344652", radius=True)
        rect(slide, 9.34, yy + 0.11, 0.18, 0.18, color, radius=True)
        text(slide, label, 9.72, yy + 0.04, 2.15, 0.3,
             size=12.3, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    footer(slide, 1, "PLANET DRUGSTORE  /  SYSTEM WALKTHROUGH  /  OCTOBER 2026")

    # 2. Overview
    slide = base_slide(prs)
    header(slide, "At a glance", "Two workflows, one shared operating rhythm",
           "Admins configure availability; employees act within the published rules; updates stay visible to both sides.")
    card(slide, 0.65, 2.15, 5.85, 3.72, "Event scheduling",
         "ADMIN / DEPARTMENT LEADS\nPrepare site details, event date and time, staffing roles, and publish the roster.\n\nEMPLOYEES\nReview published sites, claim eligible open slots, and track assignment changes.",
         TEAL, WHITE, 17, 13)
    card(slide, 6.83, 2.15, 5.85, 3.72, "Stock ordering",
         "STOCK REQUEST ADMIN\nAssign request days, set a cutoff, activate a day, and monitor incoming SR#s.\n\nEMPLOYEES\nChoose an assigned day, submit an SR#, and edit or remove it while the day is open.",
         MAROON, WHITE, 17, 13)
    rect(slide, 0.65, 6.12, 12.03, 0.55, PALE_BLUE, radius=True)
    text(slide, "Shared foundation: role-based pages  •  live Firestore data  •  optional device push notifications",
         0.88, 6.24, 11.6, 0.28, size=12, color=INK, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    footer(slide, 2)

    # 3. Event publish flow
    slide = base_slide(prs)
    header(slide, "01  /  Event scheduling", "Admin flow: prepare → publish → maintain",
           "The roster is a draft until the department lead explicitly sends or updates it.")
    xs = [0.65, 3.83, 7.01, 10.19]
    colors = [TEAL, GOLD, MAROON, RED]
    step(slide, "1", "Build the site", "Add date, time, region, event type, and location.", xs[0], 2.25, 2.55, colors[0])
    step(slide, "2", "Assign roles", "Fill Event Team and Store Reliever slots.", xs[1], 2.25, 2.55, colors[1])
    step(slide, "3", "Save or Send", "Save a draft, or Send to publish new sites to employees.", xs[2], 2.25, 2.55, colors[2])
    step(slide, "4", "Update as needed", "Changes to published rosters sync to employee views.", xs[3], 2.25, 2.55, colors[3])
    arrow(slide, 3.34, 3.15)
    arrow(slide, 6.52, 3.15)
    arrow(slide, 9.70, 3.15)
    rect(slide, 0.65, 4.82, 12.03, 1.15, MINT, radius=True)
    text(slide, "Publishing matters", 0.92, 5.04, 2.3, 0.32, size=15, color=TEAL, bold=True)
    text(slide, "Employees only see the published roster. A saved draft is still admin-side until Send / Update publishes it.",
         3.15, 5.0, 9.08, 0.55, size=13, color=INK, valign=MSO_ANCHOR.MIDDLE)
    footer(slide, 3)

    # 4. Employee event view
    slide = base_slide(prs)
    header(slide, "01  /  Event scheduling", "Employee flow: find a site and claim an eligible slot",
           "The schedule shows published sites, open roles, staffing status, and the employee’s own assignments.")
    # UI mockup
    rect(slide, 0.65, 2.05, 7.38, 4.62, WHITE, radius=True, line=LINE)
    rect(slide, 0.65, 2.05, 7.38, 0.5, INK, radius=True)
    text(slide, "EVENT SCHEDULE   /   NORTH CALOOCAN", 0.9, 2.15, 6.8, 0.25,
         size=10, color=WHITE, bold=True)
    text(slide, "Medical Mission  •  Health Center 12", 0.98, 2.88, 6.65, 0.3,
         size=16, color=INK, bold=True)
    pill(slide, "2 ROLES OPEN", 5.96, 2.84, 1.55, AMBER, "805A10", 8.5)
    text(slide, "OCT 22  •  8:00 AM", 0.98, 3.28, 3.0, 0.26,
         size=10, color=MUTED, bold=True)
    line(slide, 0.98, 3.72, 7.65, 3.72, LINE, 0.8)
    text(slide, "EVENT TEAM", 1.0, 3.92, 2.3, 0.23, size=9.5, color=TEAL, bold=True)
    text(slide, "Picker", 1.0, 4.34, 1.5, 0.28, size=12, color=INK, bold=True)
    pill(slide, "OPEN", 2.55, 4.31, 0.75, MINT, GREEN, 8)
    rect(slide, 5.6, 4.25, 1.7, 0.45, TEAL, radius=True)
    text(slide, "+ Sign me up", 5.68, 4.31, 1.54, 0.25, size=10, color=WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    line(slide, 0.98, 4.95, 7.65, 4.95, LINE, 0.8)
    text(slide, "STORE RELIEVER", 1.0, 5.13, 2.4, 0.23, size=9.5, color=MAROON, bold=True)
    text(slide, "Store coverage from assigned time onwards", 1.0, 5.51, 4.3, 0.28,
         size=11.2, color=INK)
    pill(slide, "ROLE-RESTRICTED", 5.35, 5.48, 1.75, PALE_RED, MAROON, 8)
    text(slide, "Department rules control which slots an employee can claim.",
         1.0, 6.12, 6.5, 0.3, size=10.5, color=MUTED, italic=True)
    card(slide, 8.35, 2.05, 4.33, 1.22, "1. Review", "Filter by region and inspect site details.", TEAL, WHITE, 14, 11)
    card(slide, 8.35, 3.53, 4.33, 1.22, "2. Sign up", "Claim an open role allowed for your department.", GOLD, WHITE, 14, 11)
    card(slide, 8.35, 5.01, 4.33, 1.22, "3. Track", "Assignment, cancellation, and schedule-change notices appear in the app.", MAROON, WHITE, 14, 11)
    footer(slide, 4)

    # 5. Event operations and controls
    slide = base_slide(prs)
    header(slide, "01  /  Event scheduling", "Keep the roster accurate as plans change",
           "The roster is designed for one ticket per site, clear role coverage, and visible changes.")
    columns = [
        (0.65, TEAL, "Coverage", "Sites\nUnfilled roles\nFully staffed sites", "Use the roster summary to spot staffing gaps before deployment."),
        (4.75, GOLD, "Changes", "Date / time updates\nSite and role updates\nCancel / reinstate", "Published changes flow to employees; affected people receive notices."),
        (8.85, MAROON, "Employee action", "Claim open slots\nLeave a slot\nReview notifications", "Role restrictions prevent employees from claiming incompatible slots.")
    ]
    for x, accent, title, items, note in columns:
        rect(slide, x, 2.15, 3.82, 3.72, WHITE, radius=True, line=LINE)
        rect(slide, x, 2.15, 3.82, 0.14, accent)
        text(slide, title, x + 0.25, 2.52, 3.3, 0.42, size=17, color=INK, bold=True)
        text(slide, items, x + 0.28, 3.15, 3.25, 1.37, size=13, color=INK, bold=True)
        line(slide, x + 0.28, 4.72, x + 3.48, 4.72, LINE, 0.8)
        text(slide, note, x + 0.28, 4.95, 3.25, 0.65, size=11.5, color=MUTED)
    rect(slide, 0.65, 6.12, 12.02, 0.52, PALE_BLUE, radius=True)
    text(slide, "Operational cue: check “Unfilled roles” before Send, and re-check after last-minute changes.",
         0.9, 6.23, 11.5, 0.25, size=11.8, color=INK, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    footer(slide, 5)

    # 6. Stock workflow overview
    slide = base_slide(prs)
    header(slide, "02  /  Stock ordering", "A controlled request window for each ordering day",
           "The admin opens a day and sets the cutoff; eligible employees submit SR#s before it locks.")
    step(slide, "1", "Assign days", "Admin sets which weekdays each employee may request.", 0.65, 2.28, 2.7, TEAL)
    step(slide, "2", "Open the day", "Admin confirms the date and cutoff, then activates it.", 3.77, 2.28, 2.7, GOLD)
    step(slide, "3", "Submit SR#", "Employee enters one or more numeric SR#s.", 6.89, 2.28, 2.7, MAROON)
    step(slide, "4", "Review & lock", "Admin monitors requests; cutoff closes editing.", 10.01, 2.28, 2.7, RED)
    arrow(slide, 3.43, 3.19)
    arrow(slide, 6.55, 3.19)
    arrow(slide, 9.67, 3.19)
    rect(slide, 0.65, 4.85, 12.03, 1.13, "F4E9E7", radius=True)
    text(slide, "Who can request?", 0.92, 5.08, 2.1, 0.34, size=15, color=MAROON, bold=True)
    text(slide, "Billing  •  Pharmacist  •  Pharmacy Assistant", 3.05, 5.07, 5.9, 0.35,
         size=14, color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)
    text(slide, "Branch does not change this base eligibility.", 8.95, 5.1, 3.35, 0.3,
         size=10.8, color=MUTED, italic=True, align=PP_ALIGN.RIGHT)
    footer(slide, 6)

    # 7. Stock admin setup
    slide = base_slide(prs)
    header(slide, "02  /  Stock ordering", "Admin setup: eligibility, cutoff, and activation",
           "Changes are staged and require confirmation before they are saved.")
    rect(slide, 0.65, 2.05, 7.2, 4.55, WHITE, radius=True, line=LINE)
    rect(slide, 0.65, 2.05, 7.2, 0.48, INK, radius=True)
    text(slide, "STOCK REQUEST ADMIN  /  THURSDAY", 0.92, 2.16, 6.6, 0.22,
         size=10, color=WHITE, bold=True)
    text(slide, "Employee days", 0.95, 2.85, 2.2, 0.3, size=15, color=INK, bold=True)
    text(slide, "Employee", 1.0, 3.32, 2.0, 0.23, size=9, color=MUTED, bold=True)
    text(slide, "MON   TUE   WED   THU   FRI", 3.2, 3.32, 4.15, 0.23, size=9, color=MUTED, bold=True)
    line(slide, 0.95, 3.65, 7.52, 3.65, LINE, 0.8)
    text(slide, "A. Santos", 1.0, 3.85, 1.9, 0.27, size=12, color=INK, bold=True)
    text(slide, "☑      ☐      ☑      ☑      ☐", 3.15, 3.82, 4.1, 0.3, size=12, color=TEAL, bold=True)
    text(slide, "Pending day edits", 1.0, 4.28, 2.0, 0.26, size=10.5, color=GOLD, bold=True)
    pill(slide, "CONFIRM", 5.35, 4.24, 0.95, MINT, GREEN, 8)
    pill(slide, "CANCEL", 6.4, 4.24, 0.85, PALE_RED, MAROON, 8)
    line(slide, 0.95, 4.78, 7.52, 4.78, LINE, 0.8)
    text(slide, "Cutoff", 1.0, 5.02, 1.2, 0.3, size=13, color=INK, bold=True)
    pill(slide, "2:00 PM", 2.12, 4.99, 1.05, PALE_BLUE, INK, 10)
    pill(slide, "CONFIRM CUTOFF", 3.35, 4.99, 1.7, MINT, GREEN, 8)
    pill(slide, "ACTIVATE DAY", 5.36, 4.99, 1.5, MAROON, WHITE, 8)
    text(slide, "Activation and deactivation each ask for confirmation.",
         1.0, 5.75, 6.2, 0.38, size=10.7, color=MUTED, italic=True)
    card(slide, 8.18, 2.05, 4.5, 1.18, "Unset schedule", "No requestDays field means all weekdays are allowed. An empty list means none.", GOLD, WHITE, 14, 11)
    card(slide, 8.18, 3.52, 4.5, 1.18, "Live request board", "New employee SR#s appear in the admin list and can be filtered by branch.", TEAL, WHITE, 14, 11)
    card(slide, 8.18, 4.99, 4.5, 1.18, "History & cleanup", "Completed, removed, and archived SR#s remain searchable in History.", MAROON, WHITE, 14, 11)
    footer(slide, 7)

    # 8. Stock employee flow
    slide = base_slide(prs)
    header(slide, "02  /  Stock ordering", "Employee flow: enter, update, or remove an SR#",
           "A request can be changed only while its day is active, assigned to the employee, and before cutoff.")
    rect(slide, 0.65, 2.0, 6.45, 4.72, WHITE, radius=True, line=LINE)
    rect(slide, 0.65, 2.0, 6.45, 0.5, MAROON, radius=True)
    text(slide, "EMPLOYEE STOCK REQUEST", 0.92, 2.12, 5.9, 0.24,
         size=10, color=WHITE, bold=True)
    text(slide, "Thursday, October 22", 1.0, 2.85, 4.0, 0.32, size=17, color=INK, bold=True)
    pill(slide, "OPEN UNTIL 2:00 PM", 1.0, 3.28, 1.8, MINT, GREEN, 8.5)
    text(slide, "SR#", 1.0, 3.95, 0.8, 0.24, size=10, color=MUTED, bold=True)
    rect(slide, 1.0, 4.28, 5.7, 0.56, PAPER, radius=True, line=LINE)
    text(slide, "10452, 10487", 1.2, 4.42, 4.9, 0.25, size=15, color=INK, bold=True)
    pill(slide, "UPDATE REQUEST", 1.0, 5.17, 1.7, MAROON, WHITE, 8)
    pill(slide, "REMOVE", 2.9, 5.17, 1.05, PALE_RED, RED, 8)
    text(slide, "If an SR# was already submitted, a reason is required for edits.",
         1.0, 5.92, 5.65, 0.4, size=10.5, color=MUTED, italic=True)
    card(slide, 7.48, 2.0, 5.2, 1.2, "Before cutoff", "Submit, edit, or remove your request on an active day assigned to you.", TEAL, WHITE, 14, 11.4)
    card(slide, 7.48, 3.48, 5.2, 1.2, "At cutoff", "The day locks; an existing SR# shows Completed and can no longer be edited.", GOLD, WHITE, 14, 11.4)
    card(slide, 7.48, 4.96, 5.2, 1.2, "If reopened", "When admin moves cutoff into the future, open-day SR# status returns to Pending.", MAROON, WHITE, 14, 11.4)
    footer(slide, 8)

    # 9. Cutoff scenarios
    slide = base_slide(prs)
    header(slide, "02  /  Stock ordering", "Cutoff scenarios: normal close, early close, reopen",
           "The day-level state controls whether the employee form is open; the admin’s action determines the message.")
    scenario_data = [
        (0.65, TEAL, "NORMAL CUTOFF", "2:00 PM arrives", "Day closes automatically. Existing requests become Completed; editing/removal locks."),
        (4.75, RED, "CLOSE EARLY", "Admin deactivates at 12:00 PM", "Employees see “Stock request closed early.” The day shows Closed early / Completed."),
        (8.85, GOLD, "UNCLOSE ORDER", "Admin sets a future cutoff", "Employees are told SR# entry/editing is open again. Eligible requests return to Pending.")
    ]
    for x, accent, label, title, body in scenario_data:
        rect(slide, x, 2.18, 3.82, 3.65, WHITE, radius=True, line=LINE)
        pill(slide, label, x + 0.24, 2.46, 1.65, accent, WHITE if accent != GOLD else INK, 8)
        text(slide, title, x + 0.24, 3.04, 3.25, 0.72, size=17, color=INK, bold=True)
        line(slide, x + 0.24, 3.95, x + 3.5, 3.95, LINE, 0.8)
        text(slide, body, x + 0.24, 4.18, 3.25, 1.25, size=11.7, color=MUTED)
    rect(slide, 0.65, 6.08, 12.02, 0.58, AMBER, radius=True)
    text(slide, "Push delivery: cutoff reminders are checked every 5 minutes; device push must be enabled. The in-page lock state updates live.",
         0.9, 6.21, 11.5, 0.28, size=10.8, color=INK, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    footer(slide, 9)

    # 10. Access and operating checklist
    slide = base_slide(prs)
    header(slide, "Quick reference", "Where to go and what to confirm",
           "Use the role-specific pages; each device must opt in separately for phone push notifications.")
    rect(slide, 0.65, 2.05, 5.85, 4.55, WHITE, radius=True, line=LINE)
    text(slide, "EVENT SCHEDULING", 0.95, 2.35, 4.9, 0.3, size=12, color=TEAL, bold=True)
    text(slide, "Employee schedule", 0.95, 2.9, 2.3, 0.28, size=12.5, color=INK, bold=True)
    text(slide, "/47fto0gim6", 3.6, 2.9, 2.1, 0.25, size=11, color=TEAL, bold=True)
    text(slide, "Pharmacy admin", 0.95, 3.42, 2.3, 0.28, size=12.5, color=INK, bold=True)
    text(slide, "/tz8lc2ct4y", 3.6, 3.42, 2.1, 0.25, size=11, color=TEAL, bold=True)
    text(slide, "Billing / IT admin", 0.95, 3.94, 2.3, 0.28, size=12.5, color=INK, bold=True)
    text(slide, "/jszdttqedv  •  /51qx617jve", 3.6, 3.94, 2.4, 0.38, size=10.3, color=TEAL, bold=True)
    line(slide, 0.95, 4.52, 6.15, 4.52, LINE, 0.8)
    text(slide, "Before sending", 0.95, 4.78, 2.0, 0.28, size=13, color=INK, bold=True)
    text(slide, "Check site details, role coverage, date/time, and unsaved changes.",
         0.95, 5.18, 5.15, 0.55, size=11.2, color=MUTED)

    rect(slide, 6.83, 2.05, 5.85, 4.55, WHITE, radius=True, line=LINE)
    text(slide, "STOCK REQUEST", 7.13, 2.35, 4.9, 0.3, size=12, color=MAROON, bold=True)
    text(slide, "Employee ordering", 7.13, 2.9, 2.3, 0.28, size=12.5, color=INK, bold=True)
    text(slide, "/w6n3d8ycrq", 9.75, 2.9, 2.1, 0.25, size=11, color=MAROON, bold=True)
    text(slide, "Stock request admin", 7.13, 3.42, 2.3, 0.28, size=12.5, color=INK, bold=True)
    text(slide, "/h9b4t2zvfe", 9.75, 3.42, 2.1, 0.25, size=11, color=MAROON, bold=True)
    line(slide, 7.13, 3.95, 12.33, 3.95, LINE, 0.8)
    text(slide, "Before activation", 7.13, 4.19, 2.0, 0.28, size=13, color=INK, bold=True)
    text(slide, "Confirm employee days and cutoff. An unset employee schedule allows all weekdays.",
         7.13, 4.58, 5.1, 0.52, size=11.2, color=MUTED)
    text(slide, "Before close", 7.13, 5.32, 1.6, 0.28, size=13, color=INK, bold=True)
    text(slide, "Use Unclose order only with a future cutoff; verify SR# statuses return to Pending.",
         7.13, 5.68, 5.1, 0.52, size=11.2, color=MUTED)
    footer(slide, 10)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    make_deck()